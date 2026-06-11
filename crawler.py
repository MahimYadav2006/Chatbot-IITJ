#!/usr/bin/env python3
"""
Dynamic Web Crawler for IIT Jammu Academic Departments.

Architecture:
    1. Discover URLs from validated department pages only.
    2. Normalize broken relative links against the department root.
    3. Reject generic fallback pages before expanding them.
    4. Persist a crawl manifest explaining every accept/reject decision.
"""

import argparse
import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import deque
from dataclasses import asdict, dataclass
from typing import Dict, List, Optional, Set, Tuple
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup, Comment
from playwright.sync_api import sync_playwright

from departments import (
    DEPARTMENTS,
    get_department,
    get_scraped_markdown_dir,
    resolve_department_code,
    get_scraped_data_root,
    SECTIONS,
    get_section_markdown_dir,
)
from utils import (
    MIN_CONTENT_LEN,
    PLAYWRIGHT_WAIT_MS,
    canonicalize_url,
    classify_discovered_url,
    is_binary_url,
    is_generic_content,
    is_generic_page,
)

DISCOVERY_SEED_PATHS = ("", "/", "/index.html")
CONTENT_ROOT_SELECTORS = (
    "main",
    "[role='main']",
    "article",
    ".page-content",
    ".content",
    ".entry-content",
    ".rs-services-details",
    ".rs-inner-blog",
    ".department-content",
    ".container",
)
NOISE_SELECTORS = (
    "script",
    "style",
    "noscript",
    "header",
    "footer",
    "nav",
    "form",
    "iframe[src*='youtube.com']",
    ".slick-cloned",
    ".owl-item.cloned",
    ".sr-only",
    "[hidden]",
    "[aria-hidden='true']",
)

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    from PIL import Image
except ImportError:
    Image = None


@dataclass
class LinkDecision:
    source_url: str
    target_url: str
    kind: str
    reason: str


@dataclass
class PageDecision:
    url: str
    final_url: str
    title: str
    accepted: bool
    reason: str
    text_length: int


def _compact_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()
def _score_content_node(node) -> float:
    text = _compact_text(node.get_text(" ", strip=True))
    if not text:
        return 0.0
    link_text = " ".join(a.get_text(" ", strip=True) for a in node.find_all("a"))
    return len(text) - (0.35 * len(_compact_text(link_text))) - (25 * len(node.find_all("img")))


def _select_content_root(soup: BeautifulSoup):
    best_node = None
    best_score = 0.0

    for selector in CONTENT_ROOT_SELECTORS:
        for node in soup.select(selector):
            score = _score_content_node(node)
            if score > best_score:
                best_node = node
                best_score = score

    return best_node or soup.body or soup


def _clean_soup_for_markdown(html_content: str):
    soup = BeautifulSoup(html_content, "html.parser")

    for selector in NOISE_SELECTORS:
        for tag in soup.select(selector):
            tag.decompose()

    return _select_content_root(soup)


def _extract_click_target(raw_js: str) -> Optional[str]:
    match = re.search(
        r"""(?:location\.href|window\.location(?:\.href)?|window\.open)\s*\(?\s*['"]([^'"]+)['"]""",
        raw_js or "",
        re.IGNORECASE,
    )
    if match:
        return match.group(1)
    return None


def _extract_candidate_links(soup: BeautifulSoup, base_url: str, current_url: str) -> List[str]:
    candidates: List[str] = []
    seen: Set[str] = set()

    def add_candidate(raw_target: Optional[str]):
        if not raw_target:
            return
        normalized = canonicalize_url(base_url, current_url, raw_target)
        if normalized and normalized not in seen:
            seen.add(normalized)
            candidates.append(normalized)

    for anchor in soup.find_all("a"):
        add_candidate(anchor.get("href"))
        add_candidate(anchor.get("data-href"))
        add_candidate(anchor.get("data-url"))
        add_candidate(_extract_click_target(anchor.get("onclick", "")))

    for clickable in soup.find_all(attrs={"onclick": True}):
        add_candidate(_extract_click_target(clickable.get("onclick", "")))

    return candidates


def render_page_snapshot(url: str, context) -> Dict[str, str]:
    page = context.new_page()
    try:
        print(f"Loading page via Playwright: {url}")
        navigation_succeeded = False
        try:
            page.goto(url, wait_until="networkidle", timeout=60000)
            navigation_succeeded = True
        except Exception as e:
            print(f"Playwright: networkidle failed/timed out, retrying with wait_until='load': {e}")
            try:
                page.goto(url, wait_until="load", timeout=45000)
                navigation_succeeded = True
            except Exception as e2:
                print(f"Playwright: load failed/timed out, retrying with wait_until='domcontentloaded': {e2}")
                try:
                    page.goto(url, wait_until="domcontentloaded", timeout=45000)
                    navigation_succeeded = True
                except Exception as e3:
                    print(f"Playwright: domcontentloaded also failed: {e3}")
                    # Final attempt: just commit to the URL and wait a fixed period
                    try:
                        page.goto(url, wait_until="commit", timeout=30000)
                        page.wait_for_timeout(10000)  # wait 10s for content to load
                        navigation_succeeded = True
                        print("Playwright: recovered via 'commit' + fixed wait")
                    except Exception as e4:
                        print(f"Playwright: all navigation strategies failed for {url}: {e4}")

        if navigation_succeeded:
            page.wait_for_timeout(PLAYWRIGHT_WAIT_MS)

        # Always try to extract whatever content is available
        html_content = page.content()
        final_url = page.url
        title = page.title()
        soup = BeautifulSoup(html_content, "html.parser")
        text = soup.get_text(" ", strip=True)

        if not navigation_succeeded and len(text.strip()) < 50:
            raise RuntimeError(f"All navigation strategies failed and no content loaded for {url}")

        return {
            "html": html_content,
            "final_url": final_url,
            "title": title,
            "text": text,
        }
    finally:
        page.close()


def evaluate_page(url: str, snapshot: Dict[str, str]) -> PageDecision:
    final_url = snapshot.get("final_url") or url
    title = snapshot.get("title", "")
    text = snapshot.get("text", "")
    text_length = len(_compact_text(text))

    if is_generic_page(title, text):
        return PageDecision(url, final_url, title, False, "generic-fallback-page", text_length)
    if text_length < MIN_CONTENT_LEN:
        return PageDecision(url, final_url, title, False, "too-little-text", text_length)
    return PageDecision(url, final_url, title, True, "valid-department-page", text_length)


def discover_site(base_url: str, context, max_pages: Optional[int] = None):
    queue: deque[Tuple[str, Optional[str]]] = deque()
    enqueued: Set[str] = set()
    visited: Set[str] = set()
    accepted_final_urls: Set[str] = set()
    page_snapshots: Dict[str, Dict[str, str]] = {}
    accepted_pages: List[str] = []
    binary_urls: List[str] = []
    page_decisions: List[PageDecision] = []
    link_decisions: List[LinkDecision] = []
    seen_binary_urls: Set[str] = set()

    for seed_path in DISCOVERY_SEED_PATHS:
        seed_url = canonicalize_url(base_url, base_url, base_url.rstrip("/") + seed_path)
        if seed_url and seed_url not in enqueued:
            queue.append((seed_url, None))
            enqueued.add(seed_url)

    while queue:
        current_url, _source_url = queue.popleft()
        if current_url in visited:
            continue
        if max_pages is not None and len(accepted_pages) >= max_pages:
            break

        visited.add(current_url)
        print(f"Discovering links from: {current_url}")

        try:
            snapshot = render_page_snapshot(current_url, context)
        except Exception as exc:
            page_decisions.append(
                PageDecision(
                    url=current_url,
                    final_url=current_url,
                    title="",
                    accepted=False,
                    reason=f"playwright-error: {exc}",
                    text_length=0,
                )
            )
            continue

        final_url = canonicalize_url(base_url, current_url, snapshot.get("final_url", current_url)) or current_url
        decision = evaluate_page(current_url, snapshot)
        decision = PageDecision(
            url=decision.url,
            final_url=final_url,
            title=decision.title,
            accepted=decision.accepted,
            reason=decision.reason,
            text_length=decision.text_length,
        )

        if decision.accepted and final_url in accepted_final_urls:
            decision = PageDecision(
                url=decision.url,
                final_url=final_url,
                title=decision.title,
                accepted=False,
                reason="duplicate-final-url",
                text_length=decision.text_length,
            )
        page_decisions.append(decision)
        if final_url not in page_snapshots:
            page_snapshots[final_url] = snapshot
        if not decision.accepted:
            continue

        accepted_final_urls.add(final_url)
        accepted_pages.append(final_url)

        soup = BeautifulSoup(snapshot["html"], "html.parser")
        for candidate_url in _extract_candidate_links(soup, base_url, final_url):
            kind, reason = classify_discovered_url(candidate_url, base_url, urlparse(base_url).netloc)
            link_decisions.append(
                LinkDecision(
                    source_url=final_url,
                    target_url=candidate_url,
                    kind=kind,
                    reason=reason,
                )
            )

            if kind == "page" and candidate_url not in visited and candidate_url not in enqueued:
                queue.append((candidate_url, current_url))
                enqueued.add(candidate_url)
            elif kind == "binary" and candidate_url not in seen_binary_urls:
                seen_binary_urls.add(candidate_url)
                binary_urls.append(candidate_url)

    return accepted_pages, binary_urls, page_snapshots, page_decisions, link_decisions


def html_to_markdown(element, base_url):
    """Recursively convert BeautifulSoup HTML elements into clean, relation-preserving Markdown."""
    if element is None:
        return ""

    if isinstance(element, Comment):
        return ""

    if isinstance(element, str):
        return str(element)

    name = element.name
    if not name:
        return ""

    if name in ["script", "style", "header", "footer", "nav", "noscript"]:
        return ""

    element_id = element.get("id", "") or ""
    element_classes = element.get("class", []) or []
    if element_id in ["rs-header", "rs-footer"] or any(
        css_class in ["rs-header", "rs-footer", "full-width-header"]
        for css_class in element_classes
    ):
        return ""

    if name == "img":
        alt = element.get("alt", "") or "Image"
        src = element.get("src", "")
        if src:
            abs_src = canonicalize_url(base_url, base_url, src) or src
            return f"![{alt}]({abs_src})\n\n"
        return ""

    if name == "a":
        href = element.get("href", "")
        inner_content = "".join(html_to_markdown(child, base_url) for child in element.children).strip()
        if not inner_content:
            return ""
        if href:
            abs_href = canonicalize_url(base_url, base_url, href) or href
            return f"[{inner_content}]({abs_href})"
        return inner_content

    if name in ["strong", "b"]:
        inner_content = "".join(html_to_markdown(child, base_url) for child in element.children).strip()
        return f"**{inner_content}**" if inner_content else ""

    if name in ["em", "i"]:
        inner_content = "".join(html_to_markdown(child, base_url) for child in element.children).strip()
        return f"*{inner_content}*" if inner_content else ""

    if name == "br":
        return "\n"

    if name in ["h1", "h2", "h3", "h4", "h5", "h6"]:
        level = int(name[1])
        inner_content = "".join(html_to_markdown(child, base_url) for child in element.children).strip()
        return f"\n\n{'#' * level} {inner_content}\n\n" if inner_content else ""

    if name in ["p", "div", "section", "article", "main"]:
        inner_content = "".join(html_to_markdown(child, base_url) for child in element.children).strip()
        return f"\n\n{inner_content}\n\n" if inner_content else ""

    if name in ["ul", "ol"]:
        list_items = []
        is_ordered = name == "ol"
        index = 1
        for child in element.children:
            if getattr(child, "name", None) == "li":
                item_content = "".join(html_to_markdown(grandchild, base_url) for grandchild in child.children).strip()
                if item_content:
                    prefix = f"{index}. " if is_ordered else "- "
                    list_items.append(f"{prefix}{item_content}")
                    index += 1
        return "\n\n" + "\n".join(list_items) + "\n\n" if list_items else ""

    if name == "table":
        markdown_table = []
        rows = element.find_all("tr")
        if not rows:
            return ""

        max_cols = 0
        table_data = []
        for row in rows:
            cells = row.find_all(["th", "td"])
            cell_texts = []
            for cell in cells:
                cell_text = "".join(html_to_markdown(child, base_url) for child in cell.children).strip()
                cell_text = re.sub(r"\s+", " ", cell_text)
                cell_texts.append(cell_text)
            if cell_texts:
                max_cols = max(max_cols, len(cell_texts))
                table_data.append((row.find("th") is not None, cell_texts))

        if not table_data:
            return ""

        def pad_row(row_cells, length):
            return row_cells + [""] * (length - len(row_cells))

        header_row = pad_row(table_data[0][1], max_cols)
        markdown_table.append("| " + " | ".join(header_row) + " |")
        markdown_table.append("| " + " | ".join(["---"] * max_cols) + " |")

        for _, cells in table_data[1:]:
            padded_cells = pad_row(cells, max_cols)
            markdown_table.append("| " + " | ".join(padded_cells) + " |")

        return "\n\n" + "\n".join(markdown_table) + "\n\n"

    return "".join(html_to_markdown(child, base_url) for child in element.children)


def clean_markdown(text):
    """Normalize whitespace and remove excessive newlines and duplicate adjacent lines."""
    text = re.sub(r"\n{3,}", "\n\n", text or "")
    cleaned_lines = []
    previous_nonempty = None

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            if cleaned_lines and cleaned_lines[-1] != "":
                cleaned_lines.append("")
            continue
        if line == previous_nonempty:
            continue
        cleaned_lines.append(line)
        previous_nonempty = line

    cleaned = "\n".join(cleaned_lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def parse_binary_file(response, url, output_dir):
    """Download binary files to a secure temporary path and parse their content."""
    content_type = response.headers.get("content-type", "").lower()
    ext_map = {
        "application/pdf": ".pdf",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "text/csv": ".csv",
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/gif": ".gif",
    }

    chosen_ext = None
    for mime_type, ext in ext_map.items():
        if mime_type in content_type or url.lower().endswith(ext):
            chosen_ext = ext
            break

    if not chosen_ext:
        return response.text

    with tempfile.NamedTemporaryFile(suffix=chosen_ext, dir=output_dir, delete=False) as temp_file:
        temp_path = temp_file.name
        try:
            for chunk in response.iter_content(chunk_size=8192):
                temp_file.write(chunk)
            temp_file.flush()
            temp_file.close()

            if chosen_ext == ".pdf":
                if pdfplumber is not None:
                    with pdfplumber.open(temp_path) as pdf:
                        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
                else:
                    result = subprocess.run(
                        ["pdftotext", temp_path, "-"],
                        check=True,
                        capture_output=True,
                        text=True,
                    )
                    text = result.stdout
            elif chosen_ext == ".xlsx":
                if openpyxl is None:
                    raise RuntimeError("openpyxl is not installed")
                workbook = openpyxl.load_workbook(temp_path)
                text = ""
                for sheet in workbook:
                    text += f"## Sheet: {sheet.title}\n"
                    for row in sheet.iter_rows(values_only=True):
                        text += "| " + " | ".join(str(cell) if cell else "" for cell in row) + " |\n"
                    text += "\n"
            elif chosen_ext == ".docx":
                if docx is None:
                    raise RuntimeError("python-docx is not installed")
                document = docx.Document(temp_path)
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            elif chosen_ext == ".csv":
                decoded = response.content.decode("utf-8", errors="ignore")
                reader = csv.reader(io.StringIO(decoded))
                text = "\n".join("| " + " | ".join(row) + " |" for row in reader)
            else:
                if pytesseract is None or Image is None:
                    raise RuntimeError("pytesseract/Pillow are not installed")
                text = pytesseract.image_to_string(Image.open(temp_path))

            return text
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)


def _build_output_filename(url: str, dept_code: str) -> str:
    parsed_url = urlparse(url)
    clean_path = parsed_url.path.strip("/")
    if parsed_url.query:
        clean_query = re.sub(r"[^\w\-_\.]", "_", parsed_url.query)
        clean_path = f"{clean_path}_{clean_query}"

    filename = re.sub(r"[^\w\-_\.]", "_", clean_path) or "index"
    if not filename.startswith(f"{dept_code}_"):
        filename = f"{dept_code}_{filename}"
    if not filename.endswith(".md"):
        filename = f"{filename}.md"
    return filename


def _build_quality_flags(markdown_content: str, snapshot: Optional[Dict[str, str]] = None) -> List[str]:
    flags = []
    if not markdown_content:
        flags.append("empty-markdown")
    elif len(markdown_content) < MIN_CONTENT_LEN:
        flags.append("short-markdown")

    if is_generic_content(markdown_content):
        flags.append("generic-markdown")

    if snapshot and is_generic_page(snapshot.get("title", ""), snapshot.get("text", "")):
        flags.append("generic-page-shell")

    return flags


def _build_fallback_markdown(snapshot: Optional[Dict[str, str]] = None) -> str:
    snapshot = snapshot or {}
    sections = []
    title = clean_markdown(snapshot.get("title", ""))
    visible_text = clean_markdown(snapshot.get("text", ""))

    if title:
        sections.extend(["## Page Title", "", title])

    if visible_text:
        sections.extend(["", "## Raw Visible Text", "", visible_text])

    if not sections:
        sections = ["No readable content was extracted from the rendered page."]

    return "\n".join(sections).strip()


def _html_snapshot_to_markdown(url: str, html_content: str, base_url: str, snapshot: Optional[Dict[str, str]] = None):
    content_root = _clean_soup_for_markdown(html_content)
    raw_markdown = html_to_markdown(content_root, base_url)
    markdown_content = clean_markdown(raw_markdown)
    quality_flags = _build_quality_flags(markdown_content, snapshot=snapshot)

    if not markdown_content:
        markdown_content = _build_fallback_markdown(snapshot)
        quality_flags = _build_quality_flags(markdown_content, snapshot=snapshot)

    if quality_flags:
        print(f"⚠ Warning: {', '.join(quality_flags)} for {url} – saving anyway")

    return markdown_content, quality_flags


def save_markdown_document(
    url: str,
    markdown_content: str,
    output_dir: str,
    dept_code: str,
    page_title: Optional[str] = None,
    content_flags: Optional[List[str]] = None,
):
    filename = _build_output_filename(url, dept_code)
    filepath = os.path.join(output_dir, filename)
    header_lines = [f"# Source URL: {url}"]
    if page_title:
        header_lines.append(f"# Page Title: {page_title}")
    if content_flags:
        header_lines.append(f"# Content Flags: {', '.join(content_flags)}")

    with open(filepath, "w", encoding="utf-8") as handle:
        handle.write("\n".join(header_lines) + f"\n\n{markdown_content}")
    print(f"✓ Processed: {url} -> {filename}")
    return filepath


def download_and_convert(url, output_dir, playwright_context, dept_code, base_url, cached_snapshot=None):
    """Download a URL and convert it to Markdown."""
    try:
        content_flags: List[str] = []
        page_title = None
        if is_binary_url(url):
            response = requests.get(url, stream=True, timeout=30)
            response.raise_for_status()
            markdown_content = parse_binary_file(response, url, output_dir)
            markdown_content = clean_markdown(markdown_content)
            content_flags = _build_quality_flags(markdown_content)
            if not markdown_content:
                markdown_content = "No readable content was extracted from the binary file."
                content_flags = _build_quality_flags(markdown_content)
        else:
            snapshot = cached_snapshot or render_page_snapshot(url, playwright_context)
            page_title = snapshot.get("title", "")
            markdown_content, content_flags = _html_snapshot_to_markdown(
                url,
                snapshot["html"],
                base_url,
                snapshot=snapshot,
            )

        save_markdown_document(
            url,
            markdown_content,
            output_dir,
            dept_code,
            page_title=page_title,
            content_flags=content_flags,
        )
        return markdown_content, content_flags
    except Exception as exc:
        print(f"✗ Failed: {url} – {exc}")
        return None


def _page_urls_to_export(page_decisions: List[PageDecision], page_snapshots: Dict[str, Dict[str, str]]) -> List[str]:
    urls = []
    seen = set()
    for decision in page_decisions:
        if decision.reason.startswith("playwright-error"):
            continue
        if decision.final_url in seen:
            continue
        if decision.final_url not in page_snapshots:
            continue
        seen.add(decision.final_url)
        urls.append(decision.final_url)
    return urls


def _clear_output_dir(output_dir: str):
    if not os.path.isdir(output_dir):
        return
    for name in os.listdir(output_dir):
        if name.endswith(".md") or name == "crawl_manifest.json":
            path = os.path.join(output_dir, name)
            if os.path.isfile(path):
                os.remove(path)
            else:
                shutil.rmtree(path)


def crawl_department(dept_code: str, clean_output: bool = False, max_pages: Optional[int] = None):
    canonical_code = resolve_department_code(dept_code)
    dept_config = get_department(canonical_code)
    output_dir = get_scraped_markdown_dir(canonical_code)
    os.makedirs(output_dir, exist_ok=True)

    if clean_output:
        print(f"Cleaning previous markdown output in: {output_dir}")
        _clear_output_dir(output_dir)

    base_url = dept_config["base_url"].rstrip("/")

    print("\n==================================================")
    print(f"Crawling academic department: {dept_config['full_name']} ({canonical_code.upper()})")
    print(f"Base URL: {base_url}")
    print("==================================================")

    with sync_playwright() as playwright:
        print("Phase 1: Discovering validated department pages...")
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        )

        pages, binaries, page_snapshots, page_decisions, link_decisions = discover_site(
            base_url,
            context,
            max_pages=max_pages,
        )
        page_urls_to_export = _page_urls_to_export(page_decisions, page_snapshots)
        print(
            f"Validated {len(pages)} HTML pages, rendered {len(page_urls_to_export)} HTML pages, "
            f"and discovered {len(binaries)} binary files.\n"
        )

        print("Phase 2: Converting discovered content to Markdown...")
        combined_sections = [
            f"# IIT Jammu {dept_config['full_name']} – Crawled Content",
            "",
            f"Base URL: {base_url}",
        ]

        for url in page_urls_to_export:
            result = download_and_convert(
                url,
                output_dir,
                context,
                canonical_code,
                base_url,
                cached_snapshot=page_snapshots.get(url),
            )
            if result:
                markdown, content_flags = result
                combined_sections.extend(
                    [
                        "",
                        "---",
                        "",
                        f"## {urlparse(url).path or '/'}",
                        "",
                        f"Source URL: {url}",
                        "",
                        f"Content Flags: {', '.join(content_flags) if content_flags else 'none'}",
                        "",
                        markdown,
                    ]
                )

        for url in binaries:
            result = download_and_convert(
                url,
                output_dir,
                context,
                canonical_code,
                base_url,
            )
            if result:
                markdown, content_flags = result
                combined_sections.extend(
                    [
                        "",
                        "---",
                        "",
                        f"## {urlparse(url).path or '/'}",
                        "",
                        f"Source URL: {url}",
                        "",
                        f"Content Flags: {', '.join(content_flags) if content_flags else 'none'}",
                        "",
                        markdown,
                    ]
                )

        combined_path = os.path.join(output_dir, f"00_combined_{canonical_code}_site.md")
        with open(combined_path, "w", encoding="utf-8") as handle:
            handle.write("\n".join(combined_sections).strip() + "\n")

        manifest = {
            "department_code": canonical_code,
            "department_name": dept_config["full_name"],
            "base_url": base_url,
            "accepted_pages": pages,
            "rendered_pages": page_urls_to_export,
            "accepted_binaries": binaries,
            "page_decisions": [asdict(item) for item in page_decisions],
            "link_decisions": [asdict(item) for item in link_decisions],
        }
        manifest_path = os.path.join(output_dir, "crawl_manifest.json")
        with open(manifest_path, "w", encoding="utf-8") as handle:
            json.dump(manifest, handle, indent=2)

        browser.close()

    print(f"\n✅ Completed crawling for {canonical_code.upper()}! Check '{output_dir}'.")
    print(f"Manifest written to: {os.path.join(output_dir, 'crawl_manifest.json')}")


def crawl_section(section_code: str, clean_output: bool = False, max_pages: Optional[int] = None):
    if section_code not in SECTIONS:
        raise KeyError(f"Section code '{section_code}' not found in registry.")

    section_config = SECTIONS[section_code]
    output_dir = get_section_markdown_dir(section_code)
    os.makedirs(output_dir, exist_ok=True)

    if clean_output:
        print(f"Cleaning previous markdown output in: {output_dir}")
        _clear_output_dir(output_dir)

    base_url = section_config["base_url"].rstrip("/")

    print("\n==================================================")
    print(f"Crawling section: {section_config['name']} ({section_code})")
    print(f"Base URL: {base_url}")
    print("==================================================")

    with sync_playwright() as playwright:
        print("Phase 1: Discovering validated pages...")
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        )

        pages, binaries, page_snapshots, page_decisions, link_decisions = discover_site(
            base_url,
            context,
            max_pages=max_pages,
        )
        page_urls_to_export = _page_urls_to_export(page_decisions, page_snapshots)
        print(
            f"Validated {len(pages)} HTML pages, rendered {len(page_urls_to_export)} HTML pages, "
            f"and discovered {len(binaries)} binary files.\n"
        )

        print("Phase 2: Converting discovered content to Markdown...")
        combined_sections = [
            f"# IIT Jammu {section_config['name']} – Crawled Content",
            "",
            f"Base URL: {base_url}",
        ]

        for url in page_urls_to_export:
            result = download_and_convert(
                url,
                output_dir,
                context,
                section_code,
                base_url,
                cached_snapshot=page_snapshots.get(url),
            )
            if result:
                markdown, content_flags = result
                combined_sections.extend(
                    [
                        "",
                        "---",
                        "",
                        f"## {urlparse(url).path or '/'}",
                        "",
                        f"Source URL: {url}",
                        "",
                        f"Content Flags: {', '.join(content_flags) if content_flags else 'none'}",
                        "",
                        markdown,
                    ]
                )

        for url in binaries:
            result = download_and_convert(
                url,
                output_dir,
                context,
                section_code,
                base_url,
            )
            if result:
                markdown, content_flags = result
                combined_sections.extend(
                    [
                        "",
                        "---",
                        "",
                        f"## {urlparse(url).path or '/'}",
                        "",
                        f"Source URL: {url}",
                        "",
                        f"Content Flags: {', '.join(content_flags) if content_flags else 'none'}",
                        "",
                        markdown,
                    ]
                )

        combined_path = os.path.join(output_dir, f"00_combined_{section_code}_site.md")
        with open(combined_path, "w", encoding="utf-8") as handle:
            handle.write("\n".join(combined_sections).strip() + "\n")

        manifest = {
            "section_code": section_code,
            "section_name": section_config["name"],
            "base_url": base_url,
            "accepted_pages": pages,
            "rendered_pages": page_urls_to_export,
            "accepted_binaries": binaries,
            "page_decisions": [asdict(item) for item in page_decisions],
            "link_decisions": [asdict(item) for item in link_decisions],
        }
        manifest_path = os.path.join(output_dir, "crawl_manifest.json")
        with open(manifest_path, "w", encoding="utf-8") as handle:
            json.dump(manifest, handle, indent=2)

        browser.close()

    print(f"\n✅ Completed crawling for section {section_code}! Check '{output_dir}'.")
    print(f"Manifest written to: {os.path.join(output_dir, 'crawl_manifest.json')}")


def crawl_administration(clean_output: bool = False):
    output_dir = os.path.join(get_scraped_data_root(), "administration")
    os.makedirs(output_dir, exist_ok=True)

    if clean_output:
        print(f"Cleaning previous markdown output in: {output_dir}")
        _clear_output_dir(output_dir)

    admin_urls = [
        "https://iitjammu.ac.in/board-of-governors",
        "https://iitjammu.ac.in/director",
        "https://iitjammu.ac.in/deans-and-associate-deans",
        "https://iitjammu.ac.in/registrar",
        "https://iitjammu.ac.in/finance-committee",
        "https://iitjammu.ac.in/bogchairman",
        "https://iitjammu.ac.in/building-and-works-bwc",
        "https://iitjammu.ac.in/member-senate-academic-council",
        "https://iitjammu.ac.in/annual-action--plan-committee",
    ]

    print("\n==================================================")
    print("Crawling administration pages")
    print("==================================================")

    binary_urls: List[str] = []
    seen_binary_urls: Set[str] = set()
    page_snapshots: Dict[str, Dict[str, str]] = {}
    page_decisions: List[PageDecision] = []
    link_decisions: List[LinkDecision] = []
    processed_urls: List[str] = []

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        )

        for url in admin_urls:
            print(f"Processing admin page: {url}")
            try:
                snapshot = render_page_snapshot(url, context)
            except Exception as exc:
                print(f"✗ Failed to load: {url} – {exc}")
                page_decisions.append(
                    PageDecision(
                        url=url,
                        final_url=url,
                        title="",
                        accepted=False,
                        reason=f"playwright-error: {exc}",
                        text_length=0,
                    )
                )
                continue

            final_url = snapshot.get("final_url", url)
            decision = evaluate_page(url, snapshot)
            decision = PageDecision(
                url=decision.url,
                final_url=final_url,
                title=decision.title,
                accepted=True,
                reason="explicit-administration-page",
                text_length=decision.text_length,
            )
            page_decisions.append(decision)
            page_snapshots[final_url] = snapshot
            processed_urls.append(final_url)

            # Extract links to find any binaries (like PDFs)
            soup = BeautifulSoup(snapshot["html"], "html.parser")
            base_url = "https://iitjammu.ac.in"
            for candidate_url in _extract_candidate_links(soup, base_url, final_url):
                kind, reason = classify_discovered_url(candidate_url, base_url, "iitjammu.ac.in")
                link_decisions.append(
                    LinkDecision(
                        source_url=final_url,
                        target_url=candidate_url,
                        kind=kind,
                        reason=reason,
                    )
                )
                if kind == "binary" and candidate_url not in seen_binary_urls:
                    seen_binary_urls.add(candidate_url)
                    binary_urls.append(candidate_url)

        print(f"\nRendered {len(processed_urls)} administration pages and discovered {len(binary_urls)} binary files.\n")

        print("Phase 2: Converting discovered content to Markdown...")
        combined_sections = [
            "# IIT Jammu Administration – Crawled Content",
            "",
            "Base URL: https://iitjammu.ac.in",
        ]

        for url in processed_urls:
            result = download_and_convert(
                url,
                output_dir,
                context,
                "administration",
                "https://iitjammu.ac.in",
                cached_snapshot=page_snapshots.get(url),
            )
            if result:
                markdown, content_flags = result
                combined_sections.extend(
                    [
                        "",
                        "---",
                        "",
                        f"## {urlparse(url).path or '/'}",
                        "",
                        f"Source URL: {url}",
                        "",
                        f"Content Flags: {', '.join(content_flags) if content_flags else 'none'}",
                        "",
                        markdown,
                    ]
                )

        for url in binary_urls:
            print(f"Downloading binary linked from administration: {url}")
            result = download_and_convert(
                url,
                output_dir,
                context,
                "administration",
                "https://iitjammu.ac.in",
            )
            if result:
                markdown, content_flags = result
                combined_sections.extend(
                    [
                        "",
                        "---",
                        "",
                        f"## {urlparse(url).path or '/'}",
                        "",
                        f"Source URL: {url}",
                        "",
                        f"Content Flags: {', '.join(content_flags) if content_flags else 'none'}",
                        "",
                        markdown,
                    ]
                )

        combined_path = os.path.join(output_dir, "00_combined_administration_site.md")
        with open(combined_path, "w", encoding="utf-8") as handle:
            handle.write("\n".join(combined_sections).strip() + "\n")

        manifest = {
            "department_code": "administration",
            "department_name": "Administration",
            "base_url": "https://iitjammu.ac.in",
            "accepted_pages": admin_urls,
            "rendered_pages": processed_urls,
            "accepted_binaries": binary_urls,
            "page_decisions": [asdict(item) for item in page_decisions],
            "link_decisions": [asdict(item) for item in link_decisions],
        }
        manifest_path = os.path.join(output_dir, "crawl_manifest.json")
        with open(manifest_path, "w", encoding="utf-8") as handle:
            json.dump(manifest, handle, indent=2)

    print(f"\n✅ Completed crawling for administration! Check '{output_dir}'.")
    print(f"Manifest written to: {os.path.join(output_dir, 'crawl_manifest.json')}")


def main():
    parser = argparse.ArgumentParser(description="IIT Jammu Academic Department Web Crawler")
    parser.add_argument(
        "--dept",
        default="ee",
        help="Department code to crawl (e.g. ee, cse, computer_science_engineering, bsbe, administration)",
    )
    parser.add_argument("--all", action="store_true", help="Crawl all academic departments")
    parser.add_argument(
        "--section",
        default=None,
        help="Section code to crawl (e.g. academics, alumni-affairs, cds, counselling, di, e2, saral, accounts, hindicell, ir, library, medical-centre, osd, sp, rc, sw, security, tlu, tinkerers-lab)",
    )
    parser.add_argument("--all-sections", action="store_true", help="Crawl all section websites")
    parser.add_argument(
        "--clean-output",
        action="store_true",
        help="Remove existing markdown files for the department/section before crawling",
    )
    parser.add_argument(
        "--max-pages",
        type=int,
        default=None,
        help="Optional cap for validated HTML pages to crawl (useful for debugging)",
    )
    args = parser.parse_args()

    if args.all_sections:
        print("Crawling ALL sections...")
        for section_code in SECTIONS:
            try:
                crawl_section(
                    section_code,
                    clean_output=args.clean_output,
                    max_pages=args.max_pages,
                )
            except Exception as exc:
                print(f"❌ Failed crawling section {section_code}: {exc}")
    elif args.section:
        section = args.section.lower().strip()
        try:
            crawl_section(
                section,
                clean_output=args.clean_output,
                max_pages=args.max_pages,
            )
        except Exception as exc:
            print(f"❌ Failed crawling section {section}: {exc}")
            sys.exit(1)
    elif args.all:
        print("Crawling ALL academic departments...")
        for department_code in DEPARTMENTS:
            try:
                crawl_department(
                    department_code,
                    clean_output=args.clean_output,
                    max_pages=args.max_pages,
                )
            except Exception as exc:
                print(f"❌ Failed crawling department {department_code.upper()}: {exc}")
    else:
        dept = args.dept.lower().strip()
        if dept == "administration":
            try:
                crawl_administration(clean_output=args.clean_output)
            except Exception as exc:
                print(f"❌ Failed crawling administration: {exc}")
                sys.exit(1)
        else:
            try:
                crawl_department(
                    args.dept,
                    clean_output=args.clean_output,
                    max_pages=args.max_pages,
                )
            except KeyError:
                print(
                    f"❌ Unknown department code: {args.dept}. "
                    f"Must be one of: {list(DEPARTMENTS.keys())} or 'administration'"
                )
                sys.exit(1)


if __name__ == "__main__":
    main()
